#!/usr/bin/env python3
"""
Build a biological topic-validation workbook for the GeNeDis 2026 GBM BERTopic model.

Purpose
-------
This script DOES NOT assign definitive biological labels automatically.
It prioritizes BERTopic topics for expert review using evidence from:

1. c-TF-IDF top topic terms
2. representative/top publication titles
3. MeSH descriptors attached to publications in each topic

Candidate biological families:
    - Hypoxia
    - Metabolism
    - Immune microenvironment
    - Treatment resistance
    - Therapeutic interventions

Inputs
------
bertopic_candidate_b/topic_info.csv
bertopic_candidate_b/topic_terms.csv
bertopic_candidate_b/topic_top_documents.csv
bertopic_candidate_b/topic_assignments.parquet

Outputs
-------
topic_validation/
├── topic_validation_workbook.xlsx
├── candidate_topics.csv
├── all_topic_validation_summary.csv
├── topic_family_scores.csv
├── topic_mesh_terms.csv
├── expert_annotation_template.csv
├── config.json
└── COMPLETE

Install
-------
pip install pandas pyarrow openpyxl

Run
---
python build_topic_validation.py \
    --input-dir bertopic_candidate_b \
    --output-dir topic_validation

Notes
-----
- Topic -1 (outliers) is excluded from biological topic ranking.
- Keyword families are broad screening vocabularies, not final labels.
- Higher family scores mean "review this topic first", not "this topic
  definitely belongs to that family".
"""

from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

import numpy as np
import pandas as pd


# ---------------------------------------------------------------------
# Screening vocabularies
# ---------------------------------------------------------------------
# These are deliberately broad. They support expert triage, not final
# scientific annotation.
FAMILY_KEYWORDS: Dict[str, List[str]] = {
    "hypoxia": [
        "hypoxia",
        "hypoxic",
        "hif",
        "hif1",
        "hif1a",
        "hif-1",
        "hif-1a",
        "hif-1alpha",
        "hif1alpha",
        "oxygen",
        "oxygenation",
        "anoxia",
        "anoxic",
        "pimonidazole",
        "carbonic anhydrase",
        "ca9",
        "caix",
        "vegf",
        "angiogenesis",
        "angiogenic",
    ],
    "metabolism": [
        "metabolism",
        "metabolic",
        "metabolite",
        "metabolites",
        "glycolysis",
        "glycolytic",
        "glucose",
        "lactate",
        "ldha",
        "ldh",
        "pyruvate",
        "mitochondria",
        "mitochondrial",
        "oxidative phosphorylation",
        "oxphos",
        "glutamine",
        "glutaminolysis",
        "lipid",
        "lipids",
        "fatty acid",
        "cholesterol",
        "amino acid",
        "mtor",
        "ampk",
        "autophagy",
        "ferroptosis",
        "redox",
        "ros",
        "reactive oxygen",
        "nadph",
    ],
    "immune": [
        "immune",
        "immunity",
        "immunotherapy",
        "immunosuppression",
        "immunosuppressive",
        "macrophage",
        "macrophages",
        "microglia",
        "tam",
        "tams",
        "tumor associated macrophage",
        "tumour associated macrophage",
        "myeloid",
        "t cell",
        "t cells",
        "cd8",
        "cd4",
        "treg",
        "regulatory t",
        "pd-1",
        "pd1",
        "pd-l1",
        "pdl1",
        "checkpoint",
        "checkpoints",
        "ctla4",
        "ctla-4",
        "cytokine",
        "cytokines",
        "interferon",
        "interleukin",
        "nk cell",
        "natural killer",
        "dendritic",
        "car-t",
        "cart",
        "vaccine",
        "oncolytic",
        "immune microenvironment",
    ],
    "resistance": [
        "resistance",
        "resistant",
        "radioresistance",
        "radioresistant",
        "chemoresistance",
        "chemoresistant",
        "treatment resistance",
        "therapy resistance",
        "temozolomide resistance",
        "tmz resistance",
        "recurrence",
        "recurrent",
        "relapse",
        "mgmt",
        "dna repair",
        "repair",
        "alkylating",
        "adaptive resistance",
        "drug resistance",
        "treatment failure",
        "therapy failure",
    ],
    "therapy": [
        "therapy",
        "therapeutic",
        "treatment",
        "radiotherapy",
        "radiation",
        "irradiation",
        "radiosurgery",
        "srs",
        "temozolomide",
        "tmz",
        "ttfields",
        "tumor treating fields",
        "tumour treating fields",
        "chemotherapy",
        "targeted therapy",
        "targeted therapies",
        "inhibitor",
        "inhibitors",
        "drug delivery",
        "nanoparticle",
        "nanoparticles",
        "oncolytic",
        "immunotherapy",
        "car-t",
        "vaccine",
        "surgery",
        "resection",
    ],
}


# More specific concepts that should carry more weight when detected.
HIGH_SPECIFICITY_KEYWORDS = {
    "hypoxia": {
        "hypoxia", "hypoxic", "hif", "hif1", "hif1a", "hif-1",
        "hif-1a", "hif-1alpha", "hif1alpha", "pimonidazole", "ca9", "caix",
    },
    "metabolism": {
        "glycolysis", "glycolytic", "lactate", "ldha", "pyruvate",
        "oxidative phosphorylation", "oxphos", "glutaminolysis",
        "metabolic", "metabolism",
    },
    "immune": {
        "immunosuppression", "immunosuppressive", "macrophage",
        "macrophages", "microglia", "tam", "tams", "pd-1", "pd1",
        "pd-l1", "pdl1", "checkpoint", "ctla4", "ctla-4", "car-t",
        "immune microenvironment",
    },
    "resistance": {
        "radioresistance", "radioresistant", "chemoresistance",
        "chemoresistant", "treatment resistance", "therapy resistance",
        "temozolomide resistance", "tmz resistance", "mgmt",
        "drug resistance",
    },
    "therapy": {
        "radiotherapy", "temozolomide", "ttfields", "tumor treating fields",
        "tumour treating fields", "immunotherapy", "car-t", "oncolytic",
    },
}


def parse_args():
    p = argparse.ArgumentParser(
        description="Create an expert-validation workbook for BERTopic topics."
    )

    p.add_argument(
        "--input-dir",
        type=Path,
        default=Path("bertopic_candidate_b"),
    )
    p.add_argument(
        "--output-dir",
        type=Path,
        default=Path("topic_validation"),
    )
    p.add_argument(
        "--top-mesh",
        type=int,
        default=15,
        help="Number of most frequent MeSH descriptors retained per topic.",
    )
    p.add_argument(
        "--top-titles",
        type=int,
        default=10,
        help="Number of representative/top titles shown per topic.",
    )
    p.add_argument(
        "--candidate-score-threshold",
        type=float,
        default=3.0,
        help=(
            "Minimum family evidence score for inclusion in candidate_topics.csv. "
            "This is a screening threshold, not a biological classification rule."
        ),
    )
    p.add_argument(
        "--overwrite",
        action="store_true",
    )

    return p.parse_args()


def clean_text(value) -> str:
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_text(value) -> str:
    text = clean_text(value).lower()
    # Normalize common Greek alpha variants to ASCII-ish representation.
    text = (
        text.replace("α", "alpha")
        .replace("β", "beta")
        .replace("–", "-")
        .replace("—", "-")
    )
    return text


def contains_keyword(text: str, keyword: str) -> bool:
    """
    Phrase-aware substring matching after lowercasing.

    For short alphanumeric tokens such as TAM, HIF, ROS, require
    approximate word boundaries to reduce accidental matches.
    """
    text = normalize_text(text)
    keyword = normalize_text(keyword)

    if not text or not keyword:
        return False

    if " " in keyword or "-" in keyword:
        return keyword in text

    # Simple word-like boundary. Allows biomedical forms such as hif1a.
    pattern = rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])"
    return re.search(pattern, text) is not None


def matched_keywords(text: str, family: str) -> List[str]:
    matches = [
        kw for kw in FAMILY_KEYWORDS[family]
        if contains_keyword(text, kw)
    ]
    return sorted(set(matches))


def parse_json_list(value) -> List[str]:
    if pd.isna(value) or value == "":
        return []
    if isinstance(value, list):
        return [clean_text(x) for x in value if clean_text(x)]
    try:
        parsed = json.loads(value)
        if isinstance(parsed, list):
            return [clean_text(x) for x in parsed if clean_text(x)]
    except Exception:
        pass
    return []


def load_inputs(input_dir: Path):
    paths = {
        "topic_info": input_dir / "topic_info.csv",
        "topic_terms": input_dir / "topic_terms.csv",
        "top_docs": input_dir / "topic_top_documents.csv",
        "assignments": input_dir / "topic_assignments.parquet",
    }

    for name, path in paths.items():
        if not path.exists():
            raise FileNotFoundError(f"Missing {name}: {path}")

    topic_info = pd.read_csv(paths["topic_info"])
    topic_terms = pd.read_csv(paths["topic_terms"])
    top_docs = pd.read_csv(paths["top_docs"])
    assignments = pd.read_parquet(paths["assignments"])

    # Normalize topic ids.
    for df in [topic_info, topic_terms, top_docs, assignments]:
        if "topic" in df.columns:
            df["topic"] = pd.to_numeric(df["topic"], errors="raise").astype(int)

    # BERTopic's get_topic_info normally uses capital-T Topic.
    if "Topic" in topic_info.columns:
        topic_info["topic"] = pd.to_numeric(
            topic_info["Topic"], errors="raise"
        ).astype(int)

    required_assignments = {"topic", "pmid", "publication_year"}
    missing = required_assignments - set(assignments.columns)
    if missing:
        raise ValueError(
            f"topic_assignments.parquet missing: {sorted(missing)}"
        )

    return topic_info, topic_terms, top_docs, assignments


def build_mesh_table(
    assignments: pd.DataFrame,
    top_n: int,
) -> pd.DataFrame:
    """
    Count MeSH descriptors by BERTopic topic.
    A paper contributes at most once to a given descriptor.
    """
    if "mesh_descriptors_json" not in assignments.columns:
        return pd.DataFrame(
            columns=["topic", "rank", "mesh_descriptor", "paper_count", "topic_fraction"]
        )

    rows = []

    valid = assignments[assignments["topic"] != -1].copy()

    topic_sizes = valid.groupby("topic").size().to_dict()

    for topic_id, group in valid.groupby("topic", sort=True):
        counter = Counter()

        for value in group["mesh_descriptors_json"]:
            descriptors = set(parse_json_list(value))
            counter.update(descriptors)

        denominator = topic_sizes[int(topic_id)]

        for rank, (term, count) in enumerate(
            counter.most_common(top_n),
            start=1,
        ):
            rows.append(
                {
                    "topic": int(topic_id),
                    "rank": rank,
                    "mesh_descriptor": term,
                    "paper_count": int(count),
                    "topic_fraction": float(count / denominator),
                }
            )

    return pd.DataFrame(rows)


def aggregate_topic_terms(topic_terms: pd.DataFrame) -> pd.DataFrame:
    terms = topic_terms.copy()

    terms["term"] = terms["term"].fillna("").astype(str)
    terms["ctfidf_score"] = pd.to_numeric(
        terms["ctfidf_score"], errors="coerce"
    ).fillna(0.0)

    # Text summary.
    summaries = (
        terms.sort_values(["topic", "rank"])
        .groupby("topic")
        .apply(
            lambda g: " | ".join(
                f"{row.term} ({row.ctfidf_score:.4f})"
                for row in g.itertuples()
            ),
            include_groups=False,
        )
        .rename("top_terms")
        .reset_index()
    )

    return summaries


def aggregate_titles(
    top_docs: pd.DataFrame,
    n_titles: int,
) -> pd.DataFrame:
    docs = top_docs.copy()

    if "title" not in docs.columns:
        return pd.DataFrame(columns=["topic", "representative_titles"])

    docs["title"] = docs["title"].fillna("").astype(str)

    summaries = (
        docs.sort_values(["topic", "rank"])
        .groupby("topic")
        .head(n_titles)
        .groupby("topic")
        .apply(
            lambda g: "\n".join(
                f"{int(row.rank)}. {clean_text(row.title)}"
                for row in g.itertuples()
                if clean_text(row.title)
            ),
            include_groups=False,
        )
        .rename("representative_titles")
        .reset_index()
    )

    return summaries


def aggregate_mesh(mesh_table: pd.DataFrame) -> pd.DataFrame:
    if mesh_table.empty:
        return pd.DataFrame(columns=["topic", "top_mesh"])

    return (
        mesh_table.sort_values(["topic", "rank"])
        .groupby("topic")
        .apply(
            lambda g: " | ".join(
                f"{row.mesh_descriptor} ({row.paper_count})"
                for row in g.itertuples()
            ),
            include_groups=False,
        )
        .rename("top_mesh")
        .reset_index()
    )


def topic_metadata(assignments: pd.DataFrame) -> pd.DataFrame:
    valid = assignments[assignments["topic"] != -1].copy()

    grouped = (
        valid.groupby("topic")
        .agg(
            n_documents=("pmid", "size"),
            first_year=("publication_year", "min"),
            last_year=("publication_year", "max"),
            median_year=("publication_year", "median"),
            mean_membership=(
                "cluster_membership_strength",
                "mean",
            ),
            median_membership=(
                "cluster_membership_strength",
                "median",
            ),
        )
        .reset_index()
    )

    return grouped


def build_family_scores(
    topic_terms: pd.DataFrame,
    top_docs: pd.DataFrame,
    mesh_table: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    Scoring philosophy
    ------------------
    Top terms are primary evidence.
      +2 for each family-keyword match in top topic terms
      +2 additional if the matched term is high-specificity

    Representative titles are secondary evidence.
      +1 per distinct family keyword observed across top titles
      +1 additional for high-specificity matches

    MeSH descriptors are supportive evidence.
      +1 per distinct family keyword observed in top MeSH descriptors
      +1 additional for high-specificity matches

    The score is deliberately interpretable and is only used to rank
    topics for expert review.
    """
    topics = sorted(
        set(topic_terms["topic"].unique())
        | set(top_docs["topic"].unique())
        | set(mesh_table["topic"].unique() if not mesh_table.empty else [])
    )

    score_rows = []
    evidence_rows = []

    for topic_id in topics:
        if int(topic_id) == -1:
            continue

        term_text = " ".join(
            topic_terms.loc[
                topic_terms["topic"] == topic_id,
                "term"
            ].fillna("").astype(str)
        )

        title_text = " ".join(
            top_docs.loc[
                top_docs["topic"] == topic_id,
                "title"
            ].fillna("").astype(str)
        )

        mesh_text = " ".join(
            mesh_table.loc[
                mesh_table["topic"] == topic_id,
                "mesh_descriptor"
            ].fillna("").astype(str)
        ) if not mesh_table.empty else ""

        row = {"topic": int(topic_id)}

        for family in FAMILY_KEYWORDS:
            term_matches = matched_keywords(term_text, family)
            title_matches = matched_keywords(title_text, family)
            mesh_matches = matched_keywords(mesh_text, family)

            high = HIGH_SPECIFICITY_KEYWORDS[family]

            score = 0.0

            for kw in term_matches:
                score += 2.0
                if normalize_text(kw) in {
                    normalize_text(x) for x in high
                }:
                    score += 2.0

            for kw in title_matches:
                score += 1.0
                if normalize_text(kw) in {
                    normalize_text(x) for x in high
                }:
                    score += 1.0

            for kw in mesh_matches:
                score += 1.0
                if normalize_text(kw) in {
                    normalize_text(x) for x in high
                }:
                    score += 1.0

            row[f"{family}_score"] = score
            row[f"{family}_term_matches"] = "; ".join(term_matches)
            row[f"{family}_title_matches"] = "; ".join(title_matches)
            row[f"{family}_mesh_matches"] = "; ".join(mesh_matches)

            if score > 0:
                evidence_rows.append(
                    {
                        "topic": int(topic_id),
                        "family": family,
                        "score": float(score),
                        "topic_term_matches": "; ".join(term_matches),
                        "title_matches": "; ".join(title_matches),
                        "mesh_matches": "; ".join(mesh_matches),
                    }
                )

        family_scores = {
            family: row[f"{family}_score"]
            for family in FAMILY_KEYWORDS
        }

        ordered = sorted(
            family_scores.items(),
            key=lambda x: (-x[1], x[0]),
        )

        row["primary_screen_family"] = (
            ordered[0][0] if ordered and ordered[0][1] > 0 else ""
        )
        row["primary_screen_score"] = (
            ordered[0][1] if ordered else 0.0
        )
        row["all_positive_screen_families"] = "; ".join(
            family
            for family, score in ordered
            if score > 0
        )
        row["n_positive_screen_families"] = sum(
            score > 0 for score in family_scores.values()
        )

        score_rows.append(row)

    return (
        pd.DataFrame(score_rows),
        pd.DataFrame(evidence_rows),
    )


def build_overlap_flags(scores: pd.DataFrame) -> pd.DataFrame:
    out = scores.copy()

    def positive(row, family):
        return row.get(f"{family}_score", 0) > 0

    flags = []

    for _, row in out.iterrows():
        combinations = []

        if positive(row, "hypoxia") and positive(row, "metabolism"):
            combinations.append("hypoxia+metabolism")

        if positive(row, "hypoxia") and positive(row, "immune"):
            combinations.append("hypoxia+immune")

        if positive(row, "metabolism") and positive(row, "immune"):
            combinations.append("metabolism+immune")

        if positive(row, "metabolism") and positive(row, "resistance"):
            combinations.append("metabolism+resistance")

        if positive(row, "immune") and positive(row, "resistance"):
            combinations.append("immune+resistance")

        if (
            positive(row, "hypoxia")
            and positive(row, "metabolism")
            and positive(row, "immune")
        ):
            combinations.append("hypoxia+metabolism+immune")

        if positive(row, "resistance") and positive(row, "therapy"):
            combinations.append("resistance+therapy")

        flags.append("; ".join(combinations))

    out["cross_axis_flags"] = flags
    return out


def excel_autofit_and_freeze(writer, sheet_name: str, df: pd.DataFrame):
    ws = writer.book[sheet_name]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for idx, col in enumerate(df.columns, start=1):
        series = df[col].astype(str).replace("nan", "")
        sample = series.head(250)
        max_len = max(
            [len(str(col))]
            + [len(v) for v in sample]
        )
        # Keep narrative/evidence columns readable without absurd widths.
        if any(
            key in str(col).lower()
            for key in [
                "titles",
                "terms",
                "mesh",
                "matches",
                "notes",
                "rationale",
            ]
        ):
            width = min(max(max_len + 2, 20), 60)
        else:
            width = min(max(max_len + 2, 10), 28)

        ws.column_dimensions[
            ws.cell(row=1, column=idx).column_letter
        ].width = width


def main():
    args = parse_args()
    input_dir = args.input_dir
    output_dir = args.output_dir

    if (
        output_dir.exists()
        and any(output_dir.iterdir())
        and not args.overwrite
    ):
        raise RuntimeError(
            f"Output directory '{output_dir}' is not empty. "
            "Use a new directory or --overwrite."
        )

    output_dir.mkdir(parents=True, exist_ok=True)

    topic_info, topic_terms, top_docs, assignments = load_inputs(input_dir)

    # Exclude outlier topic from expert topic validation.
    topic_terms = topic_terms[topic_terms["topic"] != -1].copy()
    top_docs = top_docs[top_docs["topic"] != -1].copy()
    assignments = assignments[assignments["topic"] != -1].copy()

    print(f"Clustered publications: {len(assignments):,}")
    print(f"Topics: {assignments['topic'].nunique():,}")

    mesh_table = build_mesh_table(
        assignments,
        top_n=args.top_mesh,
    )

    metadata = topic_metadata(assignments)
    term_summary = aggregate_topic_terms(topic_terms)
    title_summary = aggregate_titles(
        top_docs,
        n_titles=args.top_titles,
    )
    mesh_summary = aggregate_mesh(mesh_table)

    scores, evidence = build_family_scores(
        topic_terms=topic_terms,
        top_docs=top_docs,
        mesh_table=mesh_table,
    )

    scores = build_overlap_flags(scores)

    # Topic names from BERTopic.
    # load_inputs() may already have created a canonical lowercase `topic`
    # column from BERTopic's original `Topic` column. Avoid renaming `Topic`
    # again, which would create duplicate column labels.
    if "topic" not in topic_info.columns:
        if "Topic" not in topic_info.columns:
            raise ValueError(
                "topic_info.csv contains neither 'Topic' nor 'topic'."
            )
        topic_info = topic_info.copy()
        topic_info["topic"] = pd.to_numeric(
            topic_info["Topic"],
            errors="raise",
        ).astype(int)

    topic_info_small = topic_info[
        topic_info["topic"] != -1
    ].copy()

    keep_info = [
        c for c in [
            "topic",
            "Count",
            "Name",
            "Representation",
        ]
        if c in topic_info_small.columns
    ]

    # Select only the canonical lowercase topic id and discard BERTopic's
    # original capital-T Topic column if present.
    topic_info_small = (
        topic_info_small[keep_info]
        .drop_duplicates(subset=["topic"])
        .copy()
    )

    summary = (
        metadata
        .merge(topic_info_small, on="topic", how="left")
        .merge(term_summary, on="topic", how="left")
        .merge(title_summary, on="topic", how="left")
        .merge(mesh_summary, on="topic", how="left")
        .merge(scores, on="topic", how="left")
    )

    summary = summary.sort_values(
        [
            "primary_screen_score",
            "n_documents",
        ],
        ascending=[False, False],
    ).reset_index(drop=True)

    # Candidate table = any family crossing threshold.
    family_score_cols = [
        f"{family}_score"
        for family in FAMILY_KEYWORDS
    ]

    candidate_mask = (
        summary[family_score_cols]
        .max(axis=1)
        >= args.candidate_score_threshold
    )

    candidates = summary[candidate_mask].copy()

    # Expert annotation template.
    annotation_cols = [
        "topic",
        "n_documents",
        "Name",
        "primary_screen_family",
        "primary_screen_score",
        "all_positive_screen_families",
        "cross_axis_flags",
        "top_terms",
        "top_mesh",
        "representative_titles",
    ]

    annotation = candidates[
        [c for c in annotation_cols if c in candidates.columns]
    ].copy()

    annotation["expert_include"] = ""
    annotation["expert_primary_label"] = ""
    annotation["expert_secondary_labels"] = ""
    annotation["biological_coherence_1_to_5"] = ""
    annotation["relevance_to_hypoxia_metabolism_immune_resistance_1_to_5"] = ""
    annotation["expert_notes"] = ""
    annotation["final_consensus_label"] = ""

    # Save CSVs.
    summary.to_csv(
        output_dir / "all_topic_validation_summary.csv",
        index=False,
    )

    candidates.to_csv(
        output_dir / "candidate_topics.csv",
        index=False,
    )

    scores.to_csv(
        output_dir / "topic_family_scores.csv",
        index=False,
    )

    mesh_table.to_csv(
        output_dir / "topic_mesh_terms.csv",
        index=False,
    )

    evidence.to_csv(
        output_dir / "topic_family_evidence.csv",
        index=False,
    )

    annotation.to_csv(
        output_dir / "expert_annotation_template.csv",
        index=False,
    )

    # Configuration and methodological transparency.
    config = {
        "purpose": (
            "Screen BERTopic topics for expert biological review; "
            "not automatic final classification."
        ),
        "candidate_score_threshold": args.candidate_score_threshold,
        "top_mesh_per_topic": args.top_mesh,
        "top_titles_per_topic": args.top_titles,
        "family_keywords": FAMILY_KEYWORDS,
        "high_specificity_keywords": {
            k: sorted(v)
            for k, v in HIGH_SPECIFICITY_KEYWORDS.items()
        },
        "scoring": {
            "topic_term_match": 2,
            "topic_term_high_specificity_bonus": 2,
            "representative_title_match": 1,
            "representative_title_high_specificity_bonus": 1,
            "mesh_match": 1,
            "mesh_high_specificity_bonus": 1,
        },
    }

    (output_dir / "config.json").write_text(
        json.dumps(config, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    # Excel workbook.
    workbook = output_dir / "topic_validation_workbook.xlsx"

    with pd.ExcelWriter(
        workbook,
        engine="openpyxl",
    ) as writer:
        candidates.to_excel(
            writer,
            sheet_name="Candidate Topics",
            index=False,
        )
        annotation.to_excel(
            writer,
            sheet_name="Expert Annotation",
            index=False,
        )
        summary.to_excel(
            writer,
            sheet_name="All Topics",
            index=False,
        )
        scores.to_excel(
            writer,
            sheet_name="Family Scores",
            index=False,
        )
        evidence.to_excel(
            writer,
            sheet_name="Family Evidence",
            index=False,
        )
        mesh_table.to_excel(
            writer,
            sheet_name="Top MeSH",
            index=False,
        )
        topic_terms.to_excel(
            writer,
            sheet_name="Topic Terms",
            index=False,
        )
        top_docs.to_excel(
            writer,
            sheet_name="Top Documents",
            index=False,
        )

        for sheet_name, df in [
            ("Candidate Topics", candidates),
            ("Expert Annotation", annotation),
            ("All Topics", summary),
            ("Family Scores", scores),
            ("Family Evidence", evidence),
            ("Top MeSH", mesh_table),
            ("Topic Terms", topic_terms),
            ("Top Documents", top_docs),
        ]:
            excel_autofit_and_freeze(
                writer,
                sheet_name,
                df,
            )

        # Wrap text for narrative/evidence columns.
        from openpyxl.styles import Alignment, Font

        for ws in writer.book.worksheets:
            # Header emphasis.
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

            headers = {
                cell.value: cell.column
                for cell in ws[1]
            }

            wrap_headers = [
                h for h in headers
                if any(
                    key in str(h).lower()
                    for key in [
                        "titles",
                        "terms",
                        "mesh",
                        "matches",
                        "notes",
                        "labels",
                        "flags",
                    ]
                )
            ]

            for header in wrap_headers:
                col = headers[header]
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=col).alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

    (output_dir / "COMPLETE").write_text(
        "Topic validation workbook completed successfully.\n",
        encoding="utf-8",
    )

    print("\n=== TOPIC VALIDATION COMPLETE ===")
    print(f"All BERTopic topics:       {summary['topic'].nunique():,}")
    print(f"Candidate topics:          {len(candidates):,}")
    print(
        "Candidate threshold:       "
        f"{args.candidate_score_threshold:.2f}"
    )

    print("\nCandidates by primary screening family:")
    if len(candidates):
        print(
            candidates["primary_screen_family"]
            .value_counts()
            .to_string()
        )

    cross_axis = candidates[
        candidates["cross_axis_flags"].fillna("").ne("")
    ]

    print(
        f"\nCross-axis candidates:     "
        f"{len(cross_axis):,}"
    )

    print("\nTop 25 candidate topics:")
    display_cols = [
        c for c in [
            "topic",
            "n_documents",
            "Name",
            "primary_screen_family",
            "primary_screen_score",
            "all_positive_screen_families",
            "cross_axis_flags",
        ]
        if c in candidates.columns
    ]

    print(
        candidates.head(25)[display_cols].to_string(
            index=False
        )
    )

    print(f"\nWorkbook: {workbook}")
    print(f"Outputs:  {output_dir}")


if __name__ == "__main__":
    main()
